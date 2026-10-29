"""Renders the review pack README.md as an editable Word file (python-docx): headings, paragraphs, numbered and bulleted lists (one
nested level), the screens table, and inline bold, italic and code. Links keep their text.
usage: python evidence/m7-review-pack/md2docx.py IN.md OUT.docx"""
import re, sys
from docx import Document
from docx.shared import Pt

src, out = sys.argv[1], sys.argv[2]
doc = Document()
st = doc.styles['Normal']
st.font.name = 'Calibri'
st.font.size = Pt(11)
TOK = re.compile(r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|\[[^\]]+\]\([^)]+\))')


def inline(p, text):
    for part in TOK.split(text):
        if not part:
            continue
        if part.startswith('**'):
            p.add_run(part[2:-2]).bold = True
        elif part.startswith('`'):
            r = p.add_run(part[1:-1])
            r.font.name = 'Consolas'
        elif part.startswith('['):
            p.add_run(re.match(r'\[([^\]]+)\]', part).group(1))
        elif part.startswith('*'):
            p.add_run(part[1:-1]).italic = True
        else:
            p.add_run(part)


lines = open(src, encoding='utf-8').read().split('\n')
i = 0
while i < len(lines):
    ln = lines[i]
    if not ln.strip():
        i += 1
        continue
    m = re.match(r'^(#{1,3}) (.*)', ln)
    if m:
        doc.add_heading(m.group(2), level=len(m.group(1)) - 1 if len(m.group(1)) > 1 else 0)
        i += 1
        continue
    if ln.startswith('|'):
        rows = []
        while i < len(lines) and lines[i].startswith('|'):
            cells = [c.strip() for c in lines[i].strip('|').split('|')]
            if not all(re.match(r'^-+$', c) for c in cells):
                rows.append(cells)
            i += 1
        t = doc.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = 'Table Grid'
        for r, cells in enumerate(rows):
            for c, txt in enumerate(cells):
                p = t.cell(r, c).paragraphs[0]
                inline(p, txt)
                if r == 0:
                    for run in p.runs:
                        run.bold = True
        continue
    m = re.match(r'^(\s*)(\d+\.|-) (.*)', ln)
    if m:
        nested = len(m.group(1)) >= 2
        style = ('List Number' if m.group(2) != '-' else 'List Bullet') + (' 2' if nested else '')
        inline(doc.add_paragraph(style=style), m.group(3))
        i += 1
        continue
    inline(doc.add_paragraph(), ln.strip())
    i += 1
doc.save(out)
print('wrote', out)
